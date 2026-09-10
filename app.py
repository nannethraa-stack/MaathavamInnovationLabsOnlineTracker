
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from flask import Flask, flash, redirect, render_template, request, url_for, session, send_from_directory
from werkzeug.security import generate_password_hash, check_password_hash
import csv
import hashlib
import os
import unicodedata

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "project_dashboard.db"
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("MAATHAVAM_SECRET_KEY", "dev-only-change-me")


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'viewer',
        created_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS audit_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        action TEXT NOT NULL,
        entity_type TEXT NOT NULL,
        entity_id INTEGER,
        details TEXT NOT NULL DEFAULT '',
        created_at TEXT NOT NULL,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
    );

    CREATE TABLE IF NOT EXISTS projects (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        project_name TEXT NOT NULL,
        domain TEXT NOT NULL DEFAULT '',
        provisional_patent_status TEXT NOT NULL DEFAULT 'Not Filed',
        sensor TEXT NOT NULL DEFAULT 'N',
        planned_organisation TEXT NOT NULL DEFAULT '',
        status TEXT NOT NULL DEFAULT 'Planned',
        eta TEXT NOT NULL DEFAULT '',
        comments TEXT NOT NULL DEFAULT '',
        total_spend REAL NOT NULL DEFAULT 0,
        created_at TEXT NOT NULL,
        updated_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS artifacts (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        project_id INTEGER NOT NULL,
        name TEXT NOT NULL,
        description TEXT NOT NULL DEFAULT '',
        file_path TEXT,
        original_path TEXT,
        file_size INTEGER NOT NULL DEFAULT 0,
        sha256 TEXT,
        created_at TEXT NOT NULL,
        FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS comments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        project_id INTEGER NOT NULL,
        comment TEXT NOT NULL,
        created_at TEXT NOT NULL,
        FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS spend_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        project_id INTEGER NOT NULL,
        amount REAL NOT NULL,
        description TEXT NOT NULL DEFAULT '',
        spent_on TEXT NOT NULL,
        created_at TEXT NOT NULL,
        FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
    );
    """)
    # Backfill/migrate older databases safely.
    columns = {row[1] for row in conn.execute("PRAGMA table_info(artifacts)").fetchall()}
    for column, definition in (("original_path", "TEXT"), ("file_size", "INTEGER NOT NULL DEFAULT 0"), ("sha256", "TEXT")):
        if column not in columns:
            conn.execute(f"ALTER TABLE artifacts ADD COLUMN {column} {definition}")
    conn.execute("DELETE FROM projects WHERE id NOT IN (SELECT MIN(id) FROM projects GROUP BY lower(project_name))")
    conn.execute("DELETE FROM artifacts WHERE id NOT IN (SELECT MIN(id) FROM artifacts GROUP BY project_id, original_path, sha256) AND original_path IS NOT NULL AND sha256 IS NOT NULL")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_projects_name_ci ON projects(lower(project_name))")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_artifacts_identity ON artifacts(project_id, original_path, sha256) WHERE original_path IS NOT NULL AND sha256 IS NOT NULL")
    conn.commit()
    conn.close()



def audit(conn, action, entity_type, entity_id=None, details=""):
    conn.execute(
        """INSERT INTO audit_log(user_id, action, entity_type, entity_id, details, created_at)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (session.get("user_id"), action, entity_type, entity_id, details, now_iso()),
    )


def current_user():
    return session.get("username")


def require_login():
    return bool(session.get("user_id"))


def require_role(*roles):
    return session.get("role") in roles


def bootstrap_admin():
    conn = db()
    if conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
        conn.execute(
            "INSERT INTO users(username,password_hash,role,created_at) VALUES (?,?,?,?)",
            ("admin", generate_password_hash(os.environ.get("MAATHAVAM_BOOTSTRAP_PASSWORD", "admin123")), "admin", now_iso()),
        )
        conn.commit()
    conn.close()


@app.context_processor
def inject_user():
    return {"current_username": current_user(), "current_role": session.get("role")}


@app.before_request
def protect_routes():
    allowed = {"login", "static", "uploaded_file"}
    if request.endpoint not in allowed and not require_login():
        return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        conn = db()
        user = conn.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        conn.close()
        if user and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            return redirect(url_for("dashboard"))
        flash("Invalid username or password.")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/admin/users", methods=["GET", "POST"])
def users():
    if not require_role("admin"):
        return "Forbidden", 403
    conn = db()
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "viewer")
        if username and password and role in ("admin", "editor", "viewer"):
            try:
                conn.execute(
                    "INSERT INTO users(username,password_hash,role,created_at) VALUES (?,?,?,?)",
                    (username, generate_password_hash(password), role, now_iso()),
                )
                audit(conn, "CREATE", "user", None, f"Created {username} as {role}")
                conn.commit()
                flash("User created.")
            except sqlite3.IntegrityError:
                flash("Username already exists.")
        return redirect(url_for("users"))
    rows = conn.execute("SELECT id,username,role,created_at FROM users ORDER BY username").fetchall()
    conn.close()
    return render_template("users.html", users=rows)


@app.route("/admin/audit")
def audit_page():
    if not require_role("admin"):
        return "Forbidden", 403
    conn = db()
    rows = conn.execute("""
        SELECT a.*, COALESCE(u.username,'system') AS username
        FROM audit_log a LEFT JOIN users u ON u.id=a.user_id
        ORDER BY a.created_at DESC LIMIT 500
    """).fetchall()
    conn.close()
    return render_template("audit.html", rows=rows)


@app.route("/import", methods=["GET", "POST"])
def import_projects():
    if not require_role("admin", "editor"):
        return "Forbidden", 403
    if request.method == "POST":
        file = request.files.get("file")
        if not file or not file.filename:
            flash("Choose a CSV or XLSX file.")
            return redirect(url_for("import_projects"))
        name = file.filename.lower()
        records = []
        if name.endswith(".csv"):
            text = file.read().decode("utf-8-sig")
            records = list(csv.DictReader(text.splitlines()))
        elif name.endswith(".xlsx"):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(file, read_only=True, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                headers = [str(x).strip() if x is not None else "" for x in rows[0]]
                records = [dict(zip(headers, row)) for row in rows[1:]]
            except Exception as exc:
                flash(f"XLSX import failed: {exc}")
                return redirect(url_for("import_projects"))
        else:
            flash("Only CSV and XLSX are supported.")
            return redirect(url_for("import_projects"))

        conn = db()
        ts = now_iso()
        count = 0
        aliases = {
            "Project Name": "project_name",
            "Domain": "domain",
            "Provisonal Patent Filed Status": "provisional_patent_status",
            "Provisional Patent Filed Status": "provisional_patent_status",
            "sensor Y/N": "sensor",
            "Sensor Y/N": "sensor",
            "Planned Organsiation for POC": "planned_organisation",
            "Planned Organisation for POC": "planned_organisation",
            "Status": "status",
            "ETA": "eta",
            "Comments": "comments",
            "Total Spend": "total_spend",
        }
        for raw in records:
            row = {}
            for k, v in raw.items():
                if k in aliases:
                    row[aliases[k]] = "" if v is None else str(v).strip()
            if not row.get("project_name"):
                continue
            try:
                spend = float(str(row.get("total_spend", "0")).replace(",", "") or 0)
            except ValueError:
                spend = 0
            if conn.execute("SELECT 1 FROM projects WHERE lower(project_name)=lower(?)", (row.get("project_name",""),)).fetchone():
                continue
            conn.execute("""
                INSERT INTO projects(project_name,domain,provisional_patent_status,sensor,
                planned_organisation,status,eta,comments,total_spend,created_at,updated_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """, (
                row.get("project_name",""), row.get("domain",""),
                row.get("provisional_patent_status","Not Filed"), row.get("sensor","N"),
                row.get("planned_organisation",""), row.get("status","Planned") if row.get("status","Planned") in ("Planned","In Progress","POC","Validation","Blocked","Completed","On Hold") else "Planned",
                row.get("eta",""), row.get("comments",""), spend, ts, ts
            ))
            count += 1
        audit(conn, "IMPORT", "project", None, f"Imported {count} projects from {file.filename}")
        conn.commit()
        conn.close()
        flash(f"Imported {count} projects.")
        return redirect(url_for("dashboard"))
    return render_template("import.html")


@app.route("/")
def dashboard():
    sort = request.args.get("sort", "updated")
    direction = request.args.get("direction", "desc").lower()
    allowed = {
        "project": "p.project_name COLLATE NOCASE",
        "domain": "p.domain COLLATE NOCASE",
        "status": "p.status COLLATE NOCASE",
        "artifacts": "artifact_count",
        "updated": "p.updated_at",
    }
    order_col = allowed.get(sort, allowed["updated"])
    order_dir = "ASC" if direction == "asc" else "DESC"
    conn = db()
    projects = conn.execute(f"""
        SELECT p.*,
               (SELECT comment FROM comments c WHERE c.project_id=p.id ORDER BY c.created_at DESC LIMIT 1) AS latest_comment,
               (SELECT created_at FROM comments c WHERE c.project_id=p.id ORDER BY c.created_at DESC LIMIT 1) AS latest_comment_at,
               (SELECT COUNT(*) FROM artifacts a WHERE a.project_id=p.id) AS artifact_count
        FROM projects p
        ORDER BY {order_col} {order_dir}, p.id DESC
    """).fetchall()
    summary = conn.execute("""
        SELECT COUNT(*) AS concepts,
               COALESCE(SUM(CASE WHEN lower(provisional_patent_status) IN ('patented','provisional patent filed','provisional filed') THEN 1 ELSE 0 END),0) AS patented_or_filed,
               COALESCE(SUM(CASE WHEN lower(provisional_patent_status) IN ('not filed','provisional patent not filed','pending') THEN 1 ELSE 0 END),0) AS provisional_not_filed,
               COALESCE(SUM(total_spend),0) AS total_spend,
               (SELECT COUNT(*) FROM artifacts) AS artifacts
        FROM projects
    """).fetchone()
    conn.close()
    return render_template("dashboard.html", projects=projects, summary=summary, sort=sort, direction=direction)


@app.route("/project/new", methods=["GET", "POST"])
def new_project():
    if not require_role("admin", "editor"):
        return "Forbidden", 403
    if request.method == "POST":
        fields = project_fields(request)
        conn = db()
        ts = now_iso()
        if conn.execute("SELECT 1 FROM projects WHERE lower(project_name)=lower(?)", (fields[0],)).fetchone():
            conn.close()
            flash("A concept with this Project name already exists.")
            return redirect(url_for("new_project"))
        cur = conn.execute("""
            INSERT INTO projects
            (project_name, domain, provisional_patent_status, sensor,
             planned_organisation, status, eta, comments, total_spend,
             created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (*fields, ts, ts))
        project_id = cur.lastrowid

        initial_comment = request.form.get("initial_comment", "").strip()
        if initial_comment:
            conn.execute(
                "INSERT INTO comments(project_id, comment, created_at) VALUES (?, ?, ?)",
                (project_id, initial_comment, ts),
            )
        conn.commit()
        conn.close()
        flash("Project created.")
        return redirect(url_for("project", project_id=project_id))
    return render_template("project_form.html", project=None)


@app.route("/project/<int:project_id>", methods=["GET", "POST"])
def project(project_id):
    conn = db()
    p = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
    if not p:
        conn.close()
        return "Project not found", 404

    if request.method == "POST":
        if not require_role("admin", "editor"):
            return "Forbidden", 403
        action = request.form.get("action")
        ts = now_iso()

        if action == "update_project":
            fields = project_fields(request)
            conn.execute("""
                UPDATE projects SET
                    project_name=?, domain=?, provisional_patent_status=?,
                    sensor=?, planned_organisation=?, status=?, eta=?,
                    comments=?, total_spend=?, updated_at=?
                WHERE id=?
            """, (*fields, ts, project_id))
            audit(conn, "UPDATE", "project", project_id, "Project details updated")
            conn.commit()
            flash("Project details updated.")

        elif action == "add_comment":
            text = request.form.get("comment", "").strip()
            if text:
                conn.execute(
                    "INSERT INTO comments(project_id, comment, created_at) VALUES (?, ?, ?)",
                    (project_id, text, ts),
                )
                conn.execute(
                    "UPDATE projects SET comments=?, updated_at=? WHERE id=?",
                    (text, ts, project_id),
                )
                audit(conn, "CREATE", "comment", project_id, text[:200])
                conn.commit()
                flash("Comment added.")

        elif action == "add_spend":
            amount = float(request.form.get("amount", "0") or 0)
            description = request.form.get("spend_description", "").strip()
            spent_on = request.form.get("spent_on") or ts[:10]
            conn.execute("""
                INSERT INTO spend_entries(project_id, amount, description, spent_on, created_at)
                VALUES (?, ?, ?, ?, ?)
            """, (project_id, amount, description, spent_on, ts))
            total = conn.execute(
                "SELECT COALESCE(SUM(amount),0) FROM spend_entries WHERE project_id=?",
                (project_id,),
            ).fetchone()[0]
            conn.execute(
                "UPDATE projects SET total_spend=?, updated_at=? WHERE id=?",
                (total, ts, project_id),
            )
            audit(conn, "CREATE", "spend", project_id, f"{amount:.2f}")
            conn.commit()
            flash("Spend entry added.")

        elif action == "delete_artifact":
            artifact_id = int(request.form["artifact_id"])
            artifact = conn.execute("SELECT file_path FROM artifacts WHERE id=? AND project_id=?", (artifact_id, project_id)).fetchone()
            if artifact and artifact[0]:
                stored = (BASE_DIR / artifact[0]).resolve() if not artifact[0].startswith("/") else None
                if stored and UPLOAD_DIR.resolve() in stored.parents and stored.exists():
                    stored.unlink()
            conn.execute("DELETE FROM artifacts WHERE id=? AND project_id=?", (artifact_id, project_id))
            audit(conn, "DELETE", "artifact", artifact_id, "Artifact deleted")
            conn.commit()
            flash("Artifact deleted.")

        elif action == "delete_spend":
            spend_id = int(request.form["spend_id"])
            conn.execute(
                "DELETE FROM spend_entries WHERE id=? AND project_id=?",
                (spend_id, project_id),
            )
            total = conn.execute(
                "SELECT COALESCE(SUM(amount),0) FROM spend_entries WHERE project_id=?",
                (project_id,),
            ).fetchone()[0]
            conn.execute(
                "UPDATE projects SET total_spend=?, updated_at=? WHERE id=?",
                (total, ts, project_id),
            )
            audit(conn, "DELETE", "spend", spend_id, "Spend entry deleted")
            conn.commit()
            flash("Spend entry deleted.")

        return redirect(url_for("project", project_id=project_id))

    artifacts = conn.execute(
        "SELECT * FROM artifacts WHERE project_id=? ORDER BY created_at DESC",
        (project_id,),
    ).fetchall()
    comments = conn.execute(
        "SELECT * FROM comments WHERE project_id=? ORDER BY created_at DESC",
        (project_id,),
    ).fetchall()
    spend = conn.execute(
        "SELECT * FROM spend_entries WHERE project_id=? ORDER BY spent_on DESC, created_at DESC",
        (project_id,),
    ).fetchall()
    conn.close()
    return render_template(
        "project.html", project=p, artifacts=artifacts,
        comments=comments, spend=spend
    )


@app.route("/project/<int:project_id>/artifact", methods=["POST"])
def add_artifact(project_id):
    if not require_role("admin", "editor"):
        return "Forbidden", 403
    conn = db()
    if not conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone():
        conn.close()
        return "Project not found", 404

    name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip()
    file = request.files.get("file")

    if not name:
        flash("Artifact name is required.")
        conn.close()
        return redirect(url_for("project", project_id=project_id))

    saved_path = None
    if file and file.filename:
        original_name = unicodedata.normalize("NFC", Path(file.filename).name)
        safe_name = original_name.replace("/", "_").replace("\\", "_").strip() or "unnamed"
        stamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
        stored = f"{project_id}_{stamp}_{safe_name}"
        target = UPLOAD_DIR / stored
        file.save(target)
        saved_path = f"uploads/{stored}"

    conn.execute("""
        INSERT INTO artifacts(project_id, name, description, file_path, original_path, file_size, sha256, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (project_id, name, description, saved_path, name, (UPLOAD_DIR / Path(saved_path).name).stat().st_size if saved_path else 0, hashlib.sha256((UPLOAD_DIR / Path(saved_path).name).read_bytes()).hexdigest() if saved_path else None, now_iso()))
    conn.execute(
        "UPDATE projects SET updated_at=? WHERE id=?",
        (now_iso(), project_id),
    )
    audit(conn, "CREATE", "artifact", project_id, name)
    conn.commit()
    conn.close()
    flash("Artifact added.")
    return redirect(url_for("project", project_id=project_id))


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    from flask import send_from_directory
    return send_from_directory(UPLOAD_DIR, filename)


def project_fields(req):
    return (
        req.form.get("project_name", "").strip(),
        req.form.get("domain", "").strip(),
        req.form.get("provisional_patent_status", "Not Filed").strip(),
        req.form.get("sensor", "N").strip(),
        req.form.get("planned_organisation", "").strip(),
        req.form.get("status", "Planned").strip(),
        req.form.get("eta", "").strip(),
        req.form.get("comments", "").strip(),
        float(req.form.get("total_spend", "0") or 0),
    )


if __name__ == "__main__":
    init_db()
    bootstrap_admin()
    app.run(host="127.0.0.1", port=5000, debug=True)
